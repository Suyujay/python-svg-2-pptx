import pytest
from pathlib import Path
from pptx import Presentation
from surquest.utils.svg2pptx import SVG2Pptx
from tempfile import TemporaryDirectory
import os

TEST_DATA_DIR = Path("tests/data")

@pytest.mark.parametrize("json_file", [
    "01_shapes.json",
    "02_connectors.json",
    "03_text_and_groups.json",
    "04_infobox_and_icons.json",
    "05_all_features.json"
])
def test_from_json_files(json_file):
    converter = SVG2Pptx()
    input_path = TEST_DATA_DIR / json_file
    
    # Check if testing file exists
    if not input_path.exists():
        pytest.skip(f"Test data {json_file} missing.")
        
    with TemporaryDirectory() as temp_dir:
        out_pptx = Path(temp_dir) / f"{json_file}.pptx"
        
        # Test converting to pptx from file path
        converter.from_json(input_path, out_pptx)
        
        assert out_pptx.exists()
        
        # Validate pptx loads and has slides
        prs = Presentation(str(out_pptx))
        assert len(prs.slides) == 1
        
        # Also test from JSON string
        with open(input_path, "r", encoding="utf-8") as f:
            json_str = f.read()
            
        out_string_pptx = Path(temp_dir) / f"str_{json_file}.pptx"
        converter.from_json(json_str, out_string_pptx)
        assert out_string_pptx.exists()
        
        prs_str = Presentation(str(out_string_pptx))
        assert len(prs_str.slides) == 1

