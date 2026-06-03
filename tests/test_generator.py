import pytest
from pptx.util import Emu
from surquest.utils.svg2pptx.generator import PPTXBackend
from surquest.utils.svg2pptx.models import (
    IRSlide, IRRectangle, IREllipse, Geometry, Color
)

class TestPPTXBackend:
    def test_pptx_backend_empty(self):
        slide = IRSlide(width_emu=9144000, height_emu=5143500)
        backend = PPTXBackend(slide)
        prs = backend.render()
        
        assert prs.slide_width == 9144000
        assert prs.slide_height == 5143500
        assert len(prs.slides) == 1
        assert len(prs.slides[0].shapes) == 0

    def test_pptx_backend_shapes(self):
        slide = IRSlide(width_emu=9144000, height_emu=5143500)
        
        rect = IRRectangle(
            geometry=Geometry(x=Emu(100), y=Emu(100), width=Emu(200), height=Emu(200)),
            fill=Color(255, 0, 0),
            stroke=Color(0, 0, 255),
            stroke_width_emu=Emu(10),
            corner_radius_emu=0,
            shape_id="rect1"
        )
        
        ellipse = IREllipse(
            geometry=Geometry(x=Emu(400), y=Emu(100), width=Emu(200), height=Emu(200)),
            fill=Color(0, 255, 0),
            stroke=None,
            stroke_width_emu=0,
            shape_id="ellipse1"
        )
        
        slide.nodes.extend([rect, ellipse])
        
        backend = PPTXBackend(slide)
        prs = backend.render()
        
        # Render should add the shapes
        pptx_slide = prs.slides[0]
        assert len(pptx_slide.shapes) == 2
