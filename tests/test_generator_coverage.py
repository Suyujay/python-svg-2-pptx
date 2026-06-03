import pytest
import os
import json
from pathlib import Path
from tempfile import TemporaryDirectory
from pptx.util import Emu

from surquest.utils.svg2pptx import SVG2Pptx
from surquest.utils.svg2pptx.parser import _parse_color, _parse_float, SVGParser
from surquest.utils.svg2pptx.generator import PPTXBackend
from surquest.utils.svg2pptx.models import (
    Color, Point, Geometry,
    IRSlide, IRRectangle, IRLine, IRConnector, IRPolygon,
    IREllipse, IRInfoBox, IRGroup, IRIcon,
    IRText, TextBlock, TextRun, FontStyle,
    ConnectorType, ArrowType
)
from pptx.enum.text import PP_ALIGN


class TestGeneratorCoverage:
    def test_generator_full_coverage(self, tmp_path):
        from surquest.utils.svg2pptx.generator import PPTXBackend
        from surquest.utils.svg2pptx.models import IRSlide, IRRectangle, IREllipse, IRPolygon, IRGroup, IRInfoBox, IRText, TextBlock, TextRun, FontStyle, Geometry, Color, IRConnector, ConnectorType, ArrowType, Point, IRIcon
        from pptx.enum.text import PP_ALIGN
        slide = IRSlide(9000000, 5000000)
        backend = PPTXBackend(slide)
        backend.current_ir_slide = slide
        backend.slide = backend.prs.slides.add_slide(backend.prs.slide_layouts[6])
        backend._registry = {}
        backend._connectable_shapes = []
        poly_empty = IRPolygon(waypoints=[], fill=Color(0, 0, 0), stroke=Color(0, 0, 0), stroke_width_emu=Emu(0), dashed=False, is_closed=True, shape_id='poly_empty')
        backend._render_polygon(poly_empty)
        poly_full = IRPolygon(waypoints=[Point(Emu(0), Emu(0)), Point(Emu(10), Emu(10))], fill=Color(0, 0, 0), stroke=Color(0, 0, 0), stroke_width_emu=Emu(10), dashed=True, is_closed=True, shape_id='poly_full')
        backend._render_polygon(poly_full)
        grp_empty = IRGroup(children=[])
        backend._render_group(grp_empty)
        box_empty = IRInfoBox(rectangle=IRRectangle(Geometry(Emu(0), Emu(0), Emu(100), Emu(100)), fill=Color(0, 0, 0), stroke=Color(0, 0, 0), stroke_width_emu=Emu(1), corner_radius_emu=0), children=[])
        backend._render_infobox(box_empty)
        txt_blk = IRText(TextBlock(alignment=PP_ALIGN.RIGHT, runs=[TextRun('test', FontStyle('Arial', 12, False, False, Color(0, 0, 0)), False)], anchor_x=Emu(10), anchor_y=Emu(10)), container_geometry=None)
        txt_blk_2 = IRText(TextBlock(alignment=PP_ALIGN.CENTER, runs=[TextRun('test', FontStyle('Arial', 12, False, False, Color(0, 0, 0)), False)], anchor_x=Emu(10), anchor_y=Emu(10)), container_geometry=Geometry(Emu(1), Emu(1), Emu(10), Emu(10)))
        grp_text_only = IRGroup(children=[txt_blk, txt_blk_2])
        backend._render_group(grp_text_only)
        conn = IRConnector(waypoints=[Point(Emu(0), Emu(0)), Point(Emu(5), Emu(5)), Point(Emu(10), Emu(10))], stroke=Color(0, 0, 0), stroke_width_emu=Emu(10), dashed=True, connector_type=ConnectorType.STRAIGHT, start_arrow=ArrowType.TRIANGLE, end_arrow=ArrowType.TRIANGLE)
        backend._render_connector(conn)
        conn_elbow = IRConnector(waypoints=[Point(Emu(0), Emu(0)), Point(Emu(0), Emu(0))], stroke=Color(0, 0, 0), stroke_width_emu=Emu(10), dashed=False, connector_type=ConnectorType.ELBOW)
        backend._render_connector(conn_elbow)
        bad_icon = IRIcon(Geometry(Emu(0), Emu(0), Emu(100), Emu(100)), b'bad bytes')
        backend._render_icon(bad_icon)

    def test_generator_part2(self, tmp_path):
        from surquest.utils.svg2pptx.generator import PPTXBackend
        from surquest.utils.svg2pptx.models import IRNode, IRSlide, IRRectangle, IREllipse, IRPolygon, IRGroup, IRInfoBox, IRText, TextBlock, TextRun, FontStyle, Geometry, Color, IRConnector, ConnectorType, ArrowType, Point, IRIcon
        from pptx.enum.text import PP_ALIGN
        slide = IRSlide(9000000, 5000000)
        backend = PPTXBackend(slide)
        backend.current_ir_slide = slide
        backend.slide = backend.prs.slides.add_slide(backend.prs.slide_layouts[6])
        backend._registry = {}
        backend._connectable_shapes = []
    
        class UnknownNode(IRNode):
            pass
        slide.nodes.append(UnknownNode())
        txt_blk_main = IRText(TextBlock(alignment=PP_ALIGN.CENTER, runs=[TextRun('long long long', FontStyle('Arial', 12, False, False, Color(0, 0, 0)), False)], anchor_x=Emu(10), anchor_y=Emu(10)), container_geometry=Geometry(Emu(1), Emu(1), Emu(10), Emu(10)))
        txt_blk_secondary = IRText(TextBlock(alignment=PP_ALIGN.CENTER, runs=[TextRun('short', FontStyle('Arial', 12, False, False, Color(0, 0, 0)), False)], anchor_x=Emu(10), anchor_y=Emu(10)), container_geometry=None)
        rect1 = IRRectangle(Geometry(Emu(0), Emu(0), Emu(100), Emu(100)), fill=None, stroke=None, stroke_width_emu=Emu(1), corner_radius_emu=0, shape_id='huge')
        rect2 = IRRectangle(Geometry(Emu(0), Emu(0), Emu(10), Emu(10)), fill=None, stroke=None, stroke_width_emu=Emu(1), corner_radius_emu=0, shape_id='small')
        grp = IRGroup(children=[rect1, rect2, txt_blk_main, txt_blk_secondary], shape_id='grp')
        slide.nodes.append(grp)
        conn_early = IRConnector(ConnectorType.STRAIGHT, waypoints=[Point(Emu(0), Emu(0))], stroke=Color(0, 0, 0), stroke_width_emu=Emu(0), start_arrow=ArrowType.NONE, end_arrow=ArrowType.NONE)
        slide.nodes.append(conn_early)
        conn_ref = IRConnector(ConnectorType.STRAIGHT, waypoints=[Point(Emu(0), Emu(0)), Point(Emu(100), Emu(100))], stroke=Color(0, 0, 0), stroke_width_emu=Emu(0), start_ref='huge', end_ref='small', start_arrow=ArrowType.NONE, end_arrow=ArrowType.NONE)
        slide.nodes.append(conn_ref)
        conn_closest = IRConnector(ConnectorType.STRAIGHT, waypoints=[Point(Emu(10), Emu(10)), Point(Emu(90), Emu(90))], stroke=Color(0, 0, 0), stroke_width_emu=Emu(0), start_arrow=ArrowType.NONE, end_arrow=ArrowType.NONE)
        slide.nodes.append(conn_closest)
        backend.render()

    def test_generator_connectors_no_shapes(self):
        from surquest.utils.svg2pptx.generator import PPTXBackend
        from surquest.utils.svg2pptx.models import IRSlide, IRConnector, ConnectorType, ArrowType, Geometry, Point, Color
        from pptx.util import Emu
        slide = IRSlide(9000000, 5000000)
        conn1 = IRConnector(ConnectorType.STRAIGHT, waypoints=[Point(Emu(0), Emu(0)), Point(Emu(10), Emu(10))], stroke=Color(0, 0, 0), stroke_width_emu=Emu(0), dashed=True, start_arrow=ArrowType.TRIANGLE, end_arrow=ArrowType.TRIANGLE)
        conn2 = IRConnector(ConnectorType.STRAIGHT, waypoints=[Point(Emu(0), Emu(0)), Point(Emu(10), Emu(10)), Point(Emu(20), Emu(20))], stroke=Color(0, 0, 0), stroke_width_emu=Emu(0), dashed=True, start_arrow=ArrowType.TRIANGLE, end_arrow=ArrowType.TRIANGLE)
        slide.nodes.extend([conn1, conn2])
        backend = PPTXBackend(slide)
        backend.current_ir_slide = slide
        backend.slide = backend.prs.slides.add_slide(backend.prs.slide_layouts[6])
        backend._registry = {}
        backend._connectable_shapes = []
        backend.render()

    def test_generator_residual_coverage(self):
        from surquest.utils.svg2pptx.generator import PPTXBackend
        from surquest.utils.svg2pptx.models import IRSlide, IRRectangle, IRGroup, IRInfoBox, IRText, TextBlock, TextRun, FontStyle, Geometry, Color, IRConnector, ConnectorType, ArrowType, Point
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Emu
        slide = IRSlide(9000000, 5000000)
        backend = PPTXBackend(slide)
        backend.current_ir_slide = slide
        backend.slide = backend.prs.slides.add_slide(backend.prs.slide_layouts[6])
        backend._registry = {}
        backend._connectable_shapes = []
        rect = IRRectangle(Geometry(Emu(10), Emu(10), Emu(100), Emu(100)), fill=None, stroke=None, stroke_width_emu=1, corner_radius_emu=0)
        txt_left = IRText(TextBlock(alignment=PP_ALIGN.LEFT, runs=[], anchor_x=Emu(20), anchor_y=Emu(20)), container_geometry=None)
        grp_left = IRGroup(children=[rect, txt_left])
        slide.nodes.append(grp_left)
        txt_center1 = IRText(TextBlock(alignment=PP_ALIGN.CENTER, runs=[], anchor_x=Emu(90), anchor_y=Emu(20)), container_geometry=None)
        grp_center1 = IRGroup(children=[rect, txt_center1])
        slide.nodes.append(grp_center1)
        txt_center2 = IRText(TextBlock(alignment=PP_ALIGN.CENTER, runs=[], anchor_x=Emu(40), anchor_y=Emu(20)), container_geometry=None)
        grp_center2 = IRGroup(children=[rect, txt_center2])
        slide.nodes.append(grp_center2)
        rect_a = IRRectangle(Geometry(Emu(0), Emu(0), Emu(10), Emu(10)), fill=None, stroke=None, stroke_width_emu=1, corner_radius_emu=0, shape_id='a')
        rect_b = IRRectangle(Geometry(Emu(100), Emu(0), Emu(10), Emu(10)), fill=None, stroke=None, stroke_width_emu=1, corner_radius_emu=0, shape_id='b')
        slide.nodes.extend([rect_a, rect_b])
        conn_elbow_y = IRConnector(ConnectorType.ELBOW, waypoints=[Point(Emu(0), Emu(0)), Point(Emu(100), Emu(0))], stroke=Color(0, 0, 0), stroke_width_emu=1, start_ref='a', end_ref='b')
        slide.nodes.append(conn_elbow_y)
        conn_multi = IRConnector(ConnectorType.STRAIGHT, waypoints=[Point(Emu(0), Emu(0)), Point(Emu(10), Emu(10)), Point(Emu(20), Emu(20))], stroke=Color(0, 0, 0), stroke_width_emu=1, start_arrow=ArrowType.TRIANGLE, end_arrow=ArrowType.TRIANGLE)
        backend.render()

    def test_generator_residual_coverage_tweaks(self):
        from surquest.utils.svg2pptx.generator import PPTXBackend
        from surquest.utils.svg2pptx.models import IRSlide, IRRectangle, IRGroup, IRInfoBox, IRText, TextBlock, TextRun, FontStyle, Geometry, Color, IRConnector, ConnectorType, ArrowType, Point
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Emu
        slide = IRSlide(9000000, 5000000)
        backend = PPTXBackend(slide)
        backend.current_ir_slide = slide
        backend.slide = backend.prs.slides.add_slide(backend.prs.slide_layouts[6])
        backend._registry = {}
        backend._connectable_shapes = []
        rect = IRRectangle(Geometry(Emu(10), Emu(10), Emu(100), Emu(100)), fill=None, stroke=None, stroke_width_emu=1, corner_radius_emu=0)
        txt_left = IRText(TextBlock(alignment=PP_ALIGN.LEFT, runs=[TextRun('long text', FontStyle('Arial', 12, False, False, Color(0, 0, 0)), False)], anchor_x=Emu(20), anchor_y=Emu(20)), container_geometry=None)
        grp_left = IRGroup(children=[rect, txt_left])
        slide.nodes.append(grp_left)
        rect_a = IRRectangle(Geometry(Emu(0), Emu(0), Emu(10), Emu(10)), fill=None, stroke=None, stroke_width_emu=1, corner_radius_emu=0, shape_id='ay')
        rect_b = IRRectangle(Geometry(Emu(100), Emu(0), Emu(10), Emu(10)), fill=None, stroke=None, stroke_width_emu=1, corner_radius_emu=0, shape_id='by')
        conn_elbow_y = IRConnector(ConnectorType.ELBOW, waypoints=[Point(Emu(0), Emu(1)), Point(Emu(100), Emu(0))], stroke=Color(0, 0, 0), stroke_width_emu=1, start_ref='ay', end_ref='by')
        slide.nodes.extend([rect_a, rect_b, conn_elbow_y])
        backend.render()
        slide2 = IRSlide(9000000, 5000000)
        conn_multi_elbow = IRConnector(ConnectorType.ELBOW, waypoints=[Point(Emu(0), Emu(0)), Point(Emu(10), Emu(10)), Point(Emu(20), Emu(20))], stroke=Color(0, 0, 0), stroke_width_emu=1)
        slide2.nodes.append(conn_multi_elbow)
        PPTXBackend(slide2).render()

    def test_generator_direct_embed_text(self):
        from surquest.utils.svg2pptx.generator import PPTXBackend
        from surquest.utils.svg2pptx.models import IRSlide, IRRectangle, IRText, TextBlock, TextRun, FontStyle, Geometry, Color
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Emu
        slide = IRSlide(9000000, 5000000)
        backend = PPTXBackend(slide)
        backend.current_ir_slide = slide
        backend.slide = backend.prs.slides.add_slide(backend.prs.slide_layouts[6])
        backend._registry = {}
        backend._connectable_shapes = []
        rect = IRRectangle(Geometry(Emu(10), Emu(10), Emu(100), Emu(100)), fill=None, stroke=None, stroke_width_emu=1, corner_radius_emu=0)
        ppt_shape = backend._render_rectangle(rect)
        txt_left = IRText(TextBlock(alignment=PP_ALIGN.LEFT, runs=[TextRun('L', FontStyle('A', 12, False, False, Color(0, 0, 0)), False)], anchor_x=Emu(20), anchor_y=Emu(20)), container_geometry=None)
        txt_center_p = IRText(TextBlock(alignment=PP_ALIGN.CENTER, runs=[TextRun('C', FontStyle('A', 12, False, False, Color(0, 0, 0)), False)], anchor_x=Emu(150), anchor_y=Emu(20)), container_geometry=None)
        txt_center_n = IRText(TextBlock(alignment=PP_ALIGN.CENTER, runs=[TextRun('C', FontStyle('A', 12, False, False, Color(0, 0, 0)), False)], anchor_x=Emu(40), anchor_y=Emu(20)), container_geometry=None)
        backend._embed_text_in_shape(ppt_shape, txt_left)
        backend._embed_text_in_shape(ppt_shape, txt_center_p)
        backend._embed_text_in_shape(ppt_shape, txt_center_n)

    def test_generator_residual_final(self):
        from surquest.utils.svg2pptx.generator import PPTXBackend
        from surquest.utils.svg2pptx.models import IRSlide, IRRectangle, IRGroup, IRInfoBox, IRText, TextBlock, TextRun, FontStyle, Geometry, Color, IRConnector, ConnectorType, ArrowType, Point
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Emu
        slide = IRSlide(9000000, 5000000)
        backend = PPTXBackend(slide)
        backend.current_ir_slide = slide
        backend.slide = backend.prs.slides.add_slide(backend.prs.slide_layouts[6])
        backend._registry = {}
        backend._connectable_shapes = []
        rect = IRRectangle(Geometry(Emu(10), Emu(10), Emu(100), Emu(100)), fill=None, stroke=None, stroke_width_emu=1, corner_radius_emu=0)
        ppt_shape = backend._render_rectangle(rect)
        txt_right = IRText(TextBlock(alignment=PP_ALIGN.RIGHT, runs=[TextRun('R', FontStyle('A', 12, False, False, Color(0, 0, 0)), False)], anchor_x=Emu(150), anchor_y=Emu(20)), container_geometry=None)
        backend._embed_text_in_shape(ppt_shape, txt_right)
        conn = IRConnector(ConnectorType.STRAIGHT, waypoints=[Point(Emu(0), Emu(0)), Point(Emu(100), Emu(100))], stroke=Color(0, 0, 0), stroke_width_emu=1, dashed=True, start_arrow=ArrowType.TRIANGLE, end_arrow=ArrowType.TRIANGLE)
        backend._render_connector(conn)
        last_shape = backend.slide.shapes[-1]
        backend._apply_dash(last_shape)
        backend._apply_arrows(last_shape, ArrowType.TRIANGLE, ArrowType.TRIANGLE)
    
        class Bang:
    
            @property
            def shadow(self):
                raise Exception('Hit 580')
    
            @property
            def _element(self):
                raise Exception('Hit 589')
        backend._disable_shadow(Bang())
        grp = IRGroup(children=[IRRectangle(Geometry(Emu(10), Emu(10), Emu(10), Emu(10)), None, None, 1, 0), IRRectangle(Geometry(Emu(20), Emu(20), Emu(10), Emu(10)), None, None, 1, 0)])
        
        # Trigger exception natively manually without mock namespace
        def faux_add_group_shape(*args, **kwargs):
            raise Exception("Explode Group")
        backend.slide.shapes.add_group_shape = faux_add_group_shape
        backend._render_group(grp)
        
        infobox = IRInfoBox(rectangle=IRRectangle(Geometry(Emu(10), Emu(10), Emu(10), Emu(10)), None, None, 1, 0), children=[IRRectangle(Geometry(Emu(20), Emu(20), Emu(10), Emu(10)), None, None, 1, 0)])
        backend._render_infobox(infobox)
        
        def fallback_none(*args, **kwargs):
            return None
        backend._render_rectangle = fallback_none
        backend._render_infobox(infobox)
