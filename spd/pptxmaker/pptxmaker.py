"""PPTXMaker – build PowerPoint files from Canvas Document Schema JSON."""

from __future__ import annotations

import json
from typing import Any

from pptx import Presentation

from .utils import ColorUtils, UnitConverter
from .elements import ShapeHandler, SvgHandler, ImageHandler, ConnectorHandler


class PPTXMaker:
    """Creates a PowerPoint presentation from Canvas Document Schema JSON.

    Usage::

        maker = PPTXMaker(json_data)
        maker.build().save("output.pptx")

        # or from a JSON file
        PPTXMaker.from_json("input.json").build().save("output.pptx")
    """

    def __init__(self, data: dict[str, Any]) -> None:
        """
        Initialize the PPTXMaker with JSON data.

        Args:
            data: A dictionary containing the Canvas Document Schema compliant data.
        """
        self._data = data
        self._meta = data["document_metadata"]
        self._unit_converter = UnitConverter(self._meta["base_unit"])
        self._prs = Presentation()
        
        # Configure slide dimensions based on metadata
        dims = self._meta.get("dimensions")
        if dims:
            self._prs.slide_width = self._unit_converter.to_emu(dims["width"])
            self._prs.slide_height = self._unit_converter.to_emu(dims["height"])

    @classmethod
    def from_json(cls, path: str) -> PPTXMaker:
        """
        Create a PPTXMaker instance by loading data from a JSON file.

        Args:
            path: Path to the JSON file.

        Returns:
            A new PPTXMaker instance.
        """
        with open(path, encoding="utf-8") as fh:
            return cls(json.load(fh))

    # ── public API ──────────────────────────────────────────────────────

    def build(self) -> PPTXMaker:
        """
        Processes all pages in the data and builds the PowerPoint presentation.

        Returns:
            The PPTXMaker instance (allows chaining).
        """
        for page in self._data["pages"]:
            self._add_page(page)
        return self

    def save(self, path: str) -> None:
        """
        Save the generated PowerPoint presentation to the specified path.

        Args:
            path: Target file path for the .pptx file.
        """
        self._prs.save(path)

    # ── page / slide ────────────────────────────────────────────────────

    def _add_page(self, page: dict) -> None:
        """
        Adds a single slide to the presentation based on page data.

        This method follows a two-pass strategy:
        1. Add all base elements (shapes, images, SVGs) so their IDs and positions are known.
        2. Add connectors that reference those elements.
        """
        slide = self._prs.slides.add_slide(self._prs.slide_layouts[6])  # Use blank layout

        # Set slide background if specified
        bg = page.get("background")
        if bg and bg.get("type") == "solid" and bg.get("color_hex"):
            fill = slide.background.fill
            fill.solid()
            fill.fore_color.rgb = ColorUtils.hex_to_rgb(bg["color_hex"])

        # First pass: shapes and images and textboxes (track positions for connectors)
        shapes_by_id: dict[str, tuple] = {}
        for elem in page["elements"]:
            etype = elem["type"]
            handler_map = {
                "shape": (ShapeHandler, "add_shape"),
                "image": (ImageHandler, "add_image"),
                "textbox": (ShapeHandler, "add_textbox"),
                "svg_polygon": (SvgHandler, "add_svg_polygon"),
                "svg_path": (SvgHandler, "add_svg_path"),
                "svg_line": (SvgHandler, "add_svg_line"),
            }
            
            if etype in handler_map:
                handler_cls, handler_method = handler_map[etype]
                method = getattr(handler_cls, handler_method)
                shapes_by_id[elem["id"]] = (
                    method(slide, elem, self._unit_converter),
                    elem["position"],
                )

        # Second pass: connectors (they need to know about the shapes they connect to)
        for elem in page["elements"]:
            if elem["type"] == "connector":
                ConnectorHandler.add_connector(
                    slide, elem, shapes_by_id, self._unit_converter
                )
