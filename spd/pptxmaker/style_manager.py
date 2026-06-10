"""Style and text application for PPTX shapes."""

from lxml import etree
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.oxml.ns import qn
from pptx.util import Pt

from .constants import ALIGN_MAP, DASH_MAP, VANCHOR_MAP
from .utils import ColorUtils


class StyleManager:
    """
    Core logic for applying visual styles and text formatting to PowerPoint shapes.
    
    This class abstracts the low-level oxml and python-pptx calls needed to 
    match the Canvas Document Schema styles.
    """
    
    # Coefficient to adjust font sizes from SVG/pixels to Pt for better visual match
    FONT_SIZE_COEF = 0.75

    @staticmethod
    def apply_style(shape, style: dict, unit_converter) -> None:
        """
        Apply visual style properties (fill, border, opacity) to a shape.
        
        Args:
            shape: The pptx shape object.
            style: Dictionary of style properties.
            unit_converter: Unit converter for border width conversion.
        """
        # Set fill color or default to background (transparent-ish)
        if "fill_color_hex" in style:
            shape.fill.solid()
            shape.fill.fore_color.rgb = ColorUtils.hex_to_rgb(style["fill_color_hex"])
        else:
            shape.fill.background()

        # Handle opacity using low-level oxml (python-pptx doesn't have native alpha yet)
        opacity = style.get("opacity")
        if opacity is not None and opacity < 1.0:
            sp_pr = shape._element.find(qn("p:spPr"))
            if sp_pr is not None:
                solid = sp_pr.find(qn("a:solidFill"))
                if solid is not None:
                    clr = solid.find(qn("a:srgbClr"))
                    if clr is not None:
                        alpha = etree.SubElement(clr, qn("a:alpha"))
                        alpha.set("val", str(int(opacity * 100000)))

        # Handle borders/lines
        border = style.get("border")
        if border:
            ln = shape.line
            if "color_hex" in border:
                ln.color.rgb = ColorUtils.hex_to_rgb(border["color_hex"])
            if "width" in border:
                ln.width = unit_converter.to_emu(border["width"])
            if border.get("style") in DASH_MAP:
                ln.dash_style = DASH_MAP[border["style"]]
        else:
            # If no border defined, set to background (no line)
            shape.line.fill.background()

    @classmethod
    def apply_text(cls, shape, tc: dict, unit_converter) -> None:
        """
        Apply text content, alignment, and character formatting to a shape.
        
        Args:
            shape: The pptx shape object.
            tc: Text content dictionary from the schema.
            unit_converter: Unit converter for margin conversion.
        """
        if not hasattr(shape, "text_frame") or not tc:
            return

        tf = shape.text_frame
        tf.word_wrap = tc.get("word_wrap", True)
        
        # Reset default inner margins for better alignment control
        tf.margin_left = 0
        tf.margin_right = 0
        tf.margin_top = 0
        tf.margin_bottom = 0
        
        margins = tc.get("margins", {})
        if "top" in margins:
            tf.margin_top = unit_converter.to_emu(margins["top"])
        if "bottom" in margins:
            tf.margin_bottom = unit_converter.to_emu(margins["bottom"])
        if "left" in margins:
            tf.margin_left = unit_converter.to_emu(margins["left"])
        if "right" in margins:
            tf.margin_right = unit_converter.to_emu(margins["right"])

        # Ensure text fits within the shape boundaries
        tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE

        # Apply vertical alignment via body properties
        anchor = VANCHOR_MAP.get(tc.get("vertical_alignment", "MIDDLE"))
        if anchor:
            body_pr = tf._txBody.find(qn("a:bodyPr"))
            if body_pr is not None:
                body_pr.set("anchor", anchor)

        # Process paragraphs and runs
        for i, pdata in enumerate(tc.get("paragraphs", [])):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            
            align = pdata.get("alignment")
            if align in ALIGN_MAP:
                para.alignment = ALIGN_MAP[align]
                
            for rdata in pdata.get("runs", []):
                run = para.add_run()
                run.text = rdata["text"]
                f = run.font
                
                # Apply character formatting
                if "font_name" in rdata:
                    f.name = rdata["font_name"]
                if "font_size" in rdata:
                    f.size = Pt(rdata["font_size"] * cls.FONT_SIZE_COEF)
                if "bold" in rdata:
                    f.bold = rdata["bold"]
                if "italic" in rdata:
                    f.italic = rdata["italic"]
                if "underline" in rdata:
                    f.underline = rdata["underline"]
                if "color_hex" in rdata:
                    f.color.rgb = ColorUtils.hex_to_rgb(rdata["color_hex"])
